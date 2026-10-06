"""Versioned XLSX question bank and server-issued, immutable assessment forms."""
import copy
import hashlib
import io
import json
import math
import random
import re
import time
import uuid
import zipfile
from pathlib import Path

from .errors import ApiError

DIMENSIONS = 'CRDGAK'
WEIGHTS = {'anchor':1.0, 'standard':1.0, 'challenge':1.0}
HEADERS = ['题目ID','题池','维度','难度','场景','题干','O1文本','O1分','O2文本','O2分','O3文本','O3分','O4文本','O4分','来源ID','公开署名','来源URL','前置信息','facet','状态','O1新颖度','O2新颖度','O3新颖度','O4新颖度','计分权重']
MAX_BYTES = 4*1024*1024
VERSION_PREFIX = 'cloud-bank-'
ID_RE = re.compile(r'^[A-Za-z0-9][A-Za-z0-9_.-]{0,79}$')
QUESTION_REFRESH_LIMIT = 6


def encode(value):
    return json.dumps(value,ensure_ascii=False,separators=(',',':'),sort_keys=True)


class BankValidationError(Exception):
    def __init__(self, errors):
        self.errors=errors[:100]


def parse_xlsx(content):
    """Never evaluate formulas or open external workbook links."""
    if len(content)>MAX_BYTES: raise BankValidationError(['文件最多 4MB'])
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            entries=z.infolist()
            if len(entries)>500 or sum(e.file_size for e in entries)>24*1024*1024:
                raise ValueError('解压后内容超出限制')
            if any('vbaProject' in e.filename for e in entries): raise ValueError('不支持宏文件')
        import openpyxl
        book=openpyxl.load_workbook(io.BytesIO(content),read_only=True,data_only=False,keep_links=False)
    except Exception as error:
        raise BankValidationError(['无法读取 XLSX 文件，请使用后台模板（不支持宏/加密文件）']) from error
    errors=[];warnings=[];records=[];source_records={}
    try:
        if '减压题库' in book.sheetnames:
            if any(name in book.sheetnames for name in ('题库','轻松题库','挑战题库')):
                raise BankValidationError(['计分题与减压题请分开导入，避免启用范围不明确'])
            from .relaxation_bank import parse_relaxation
            return parse_relaxation(book)
        has_ledger='来源与版权台账' in book.sheetnames
        if has_ledger:
            sheet=book['来源与版权台账']; rows=sheet.iter_rows(min_row=4,max_row=5005,max_col=30,values_only=True)
            headers=list(next(rows,[]))
            for number,values in enumerate(rows,5):
                if not any(x is not None for x in values):continue
                row=dict(zip(headers,values)); sid=str(row.get('来源ID') or '').strip()
                if not sid:errors.append(f'来源与版权台账 第 {number} 行：缺少来源ID');continue
                if sid in source_records:errors.append(f'来源与版权台账 第 {number} 行：来源ID重复 {sid}')
                source_records[sid]={str(k):str(v or '') for k,v in row.items() if k}
        sheet_names=['题库'] if '题库' in book.sheetnames else [n for n in ('轻松题库','挑战题库') if n in book.sheetnames]
        if not sheet_names:raise BankValidationError(['缺少“题库”工作表，或“轻松题库/挑战题库”工作表'])
        seen=set()
        from .paid_contract import contract
        baseline={q['id']:q for q in contract()['assessmentVersions']['candidate-full-1.0.0']['items']}
        for name in sheet_names:
            first=1 if name=='题库' else 4
            book[name].reset_dimensions()
            rows=book[name].iter_rows(min_row=first,max_col=50,values_only=True)
            headers=list(next(rows,[]))
            required=HEADERS[:15]
            if any(h not in headers for h in required):
                errors.append(name+'：缺少字段 '+', '.join(h for h in required if h not in headers));continue
            if len([h for h in headers if h])!=len(set(h for h in headers if h)):
                errors.append(name+'：存在重复表头');continue
            for number,values in enumerate(rows,first+1):
                if number>2006:
                    errors.append(name+'：行数超出限制，请删除多余空行');break
                if not any(v is not None for v in values):continue
                row=dict(zip(headers,values)); loc=f'{name} 第 {number} 行'
                try:
                    def text(key, limit=1200, required=False):
                        value=str(row.get(key) or '').strip()
                        if required and not value:raise ValueError(key+' 不能为空')
                        if len(value)>limit:raise ValueError(key+f' 最多 {limit} 字符')
                        if value.startswith('='):raise ValueError(key+' 必须填值，不能填公式')
                        return value
                    qid=text('题目ID',80,True)
                    if not ID_RE.fullmatch(qid) or qid.startswith('ICE-'):raise ValueError('题目ID只允许字母数字 ._-，不能使用 ICE- 前缀')
                    if qid in seen:raise ValueError('题目ID重复 '+qid)
                    seen.add(qid)
                    pool=text('题池',20,True);dim=text('维度',1,True)
                    if pool not in WEIGHTS or dim not in DIMENSIONS:raise ValueError('题池或维度不在允许范围')
                    mode={'anchor':'shared','standard':'standard','challenge':'challenge'}[pool]
                    if row.get('模式') and text('模式')!=mode:raise ValueError('模式与题池不一致')
                    difficulty=row.get('难度')
                    if type(difficulty) not in (int,float) or difficulty not in (1,2,3):raise ValueError('难度必须为数字 1/2/3')
                    options=[]
                    for i in range(1,5):
                        label=text(f'O{i}文本',600,True); score=row.get(f'O{i}分')
                        if type(score) not in (int,float) or score not in (0,1,2,3):raise ValueError(f'O{i}分必须为数字 0/1/2/3')
                        previous=next((o for o in baseline.get(qid,{}).get('options',[]) if o['label']==label),None)
                        style=copy.deepcopy(previous['style']) if previous else {'novelty':1.5,'warmth':1.5,'directness':1.5}
                        novelty=row.get(f'O{i}新颖度')
                        if novelty is not None:
                            if type(novelty) not in (int,float) or not math.isfinite(novelty) or not 0<=novelty<=3:raise ValueError(f'O{i}新颖度为 0 至 3 数值')
                            style['novelty']=novelty
                        options.append({'id':f'O{i}','label':label,'score':int(score),'style':style})
                    if sorted(o['score'] for o in options)!=[0,1,2,3]:raise ValueError('每题四个分值须分别为 0、1、2、3')
                    if len({o['label'] for o in options})!=4:raise ValueError('选项文本重复')
                    source_id=text('来源ID',80,True)
                    if not ID_RE.fullmatch(source_id):raise ValueError('来源ID格式不正确')
                    source=source_records.get(source_id)
                    if has_ledger and not source:raise ValueError('来源ID未在来源与版权台账中找到')
                    label=text('公开署名',200) or (source or {}).get('公开署名','')
                    url=text('来源URL',1000) or (source or {}).get('官方URL','')
                    if not label:raise ValueError('缺少公开署名（可填原创题）')
                    if url and not re.match(r'^(https?://|llmer://)',url):raise ValueError('来源URL必须以 http(s):// 或 llmer:// 开头')
                    prelude=text('前置信息') or text('连续前屏信息')
                    if text('连续性')=='是' and not prelude:raise ValueError('连续性题必须提供前屏信息')
                    prompt=text('题干',1200,True)
                    if prelude:
                        prompt=prelude+'\n\n'+prompt
                        warnings.append(loc+'：前置信息合并为独立题，不当作跨题记忆证据')
                    status=text('状态',40) or 'candidate/draft'
                    if status not in ('candidate/draft','reviewed','pilot','active','retired'):raise ValueError('状态不在允许范围')
                    previous=baseline.get(qid,{})
                    weight=row.get('计分权重')
                    if weight is None:
                        weight=previous.get('scoringWeight',1.0) if previous.get('pool')==pool and previous.get('dimension')==dim else 1.0
                    if type(weight) not in (int,float) or not math.isfinite(weight) or not 0<weight<=3:
                        raise ValueError('计分权重须大于 0 且不超过 3')
                    records.append({'id':qid,'pool':pool,'mode':mode,'dimension':dim,'difficulty':int(difficulty),
                        'scene':text('场景',120,True),'prompt':prompt,'options':options,'facet':text('facet',100) or 'general',
                        'scoringWeight':weight,'continuityRequired':False,'crossQuestionDependency':False,
                        'status':status,'sourceRefId':source_id,'sourceLabel':label,
                        'sourceCategory':text('来源类别',80) or (source or {}).get('类别','original')})
                    if source is None: source_records[source_id]={'来源ID':source_id,'公开署名':label,'官方URL':url}
                except ValueError as error:errors.append(loc+'：'+str(error))
        if not records:errors.append('没有有效题目')
        if len(records)>2000:errors.append('一次最多 2000 题')
        # Disabled/retired rows are preserved in a version but never drawn.
        active=[q for q in records if q['status']!='retired']
        for d in DIMENSIONS:
            for pool, minimum in [('anchor',2),('standard',1),('challenge',1)]:
                n=sum(q['dimension']==d and q['pool']==pool for q in active)
                if n<minimum:errors.append(f'{d}/{pool} 至少 {minimum} 题，当前 {n} 题')
        if errors:raise BankValidationError(errors)
        return {'questions':records,'sources':list(source_records.values()),'warnings':warnings[:100]}
    finally:book.close()


class QuestionBankService:
    def __init__(self,db):
        self.db=db
        with db.connect() as c:
            c.executescript('''
                CREATE TABLE IF NOT EXISTS question_bank_versions(id TEXT PRIMARY KEY,name TEXT NOT NULL,content_hash TEXT NOT NULL UNIQUE,source_hash TEXT NOT NULL,payload_json TEXT NOT NULL,created_ms INTEGER NOT NULL,created_by TEXT NOT NULL);
                CREATE TABLE IF NOT EXISTS question_bank_active(singleton INTEGER PRIMARY KEY CHECK(singleton=1),version_id TEXT NOT NULL REFERENCES question_bank_versions(id));
                CREATE TABLE IF NOT EXISTS question_bank_forms(id TEXT PRIMARY KEY,user_id TEXT NOT NULL REFERENCES users(id) ON DELETE CASCADE,request_id TEXT NOT NULL,version_id TEXT NOT NULL REFERENCES question_bank_versions(id),mode TEXT NOT NULL,question_ids_json TEXT NOT NULL,created_ms INTEGER NOT NULL,UNIQUE(user_id,request_id));
                CREATE TABLE IF NOT EXISTS question_bank_results(form_id TEXT PRIMARY KEY REFERENCES question_bank_forms(id) ON DELETE CASCADE,snapshot_json TEXT NOT NULL,submitted_ms INTEGER NOT NULL);
                CREATE INDEX IF NOT EXISTS question_bank_forms_created ON question_bank_forms(created_ms);
                CREATE INDEX IF NOT EXISTS question_bank_forms_user_mode ON question_bank_forms(user_id,mode,created_ms);
                CREATE TABLE IF NOT EXISTS question_bank_compositions(form_id TEXT PRIMARY KEY REFERENCES question_bank_forms(id) ON DELETE CASCADE,payload_json TEXT NOT NULL);
                CREATE TABLE IF NOT EXISTS relaxation_bank_active(singleton INTEGER PRIMARY KEY CHECK(singleton=1),version_id TEXT NOT NULL REFERENCES question_bank_versions(id));
                CREATE TABLE IF NOT EXISTS question_bank_assignments(
                    user_id TEXT NOT NULL REFERENCES users(id) ON DELETE CASCADE,
                    mode TEXT NOT NULL CHECK(mode IN ('standard','challenge')),
                    source_form_id TEXT NOT NULL REFERENCES question_bank_forms(id) ON DELETE CASCADE,
                    refresh_used INTEGER NOT NULL DEFAULT 0 CHECK(refresh_used>=0),
                    PRIMARY KEY(user_id,mode));
            ''')
            columns={r['name'] for r in c.execute('PRAGMA table_info(question_bank_forms)')}
            if 'include_relaxation' not in columns:
                c.execute('ALTER TABLE question_bank_forms ADD COLUMN include_relaxation INTEGER NOT NULL DEFAULT 0')
            if 'shuffle_relaxation' not in columns:
                c.execute('ALTER TABLE question_bank_forms ADD COLUMN shuffle_relaxation INTEGER NOT NULL DEFAULT 0')
            if 'source_form_id' not in columns:
                c.execute('ALTER TABLE question_bank_forms ADD COLUMN source_form_id TEXT')
            if 'refresh_requested' not in columns:
                c.execute('ALTER TABLE question_bank_forms ADD COLUMN refresh_requested INTEGER NOT NULL DEFAULT 0')

    def import_xlsx(self,content,name,author):
        return self._import_payload(parse_xlsx(content), content, name, author)

    def import_json(self, content, name, author):
        from .question_bank_json import parse_json
        return self._import_payload(parse_json(content), content, name, author)

    def _import_payload(self, payload, content, name, author):
        raw=encode(payload); fingerprint=hashlib.sha256(raw.encode()).hexdigest()
        is_relaxation=payload.get('kind')=='relaxation'
        vid=('relax-bank-' if is_relaxation else VERSION_PREFIX)+fingerprint[:24]
        active_table='relaxation_bank_active' if is_relaxation else 'question_bank_active'
        with self.db.connect() as c:
            current=c.execute(f'SELECT v.payload_json FROM {active_table} a JOIN question_bank_versions v ON a.version_id=v.id').fetchone()
            old={q['id']:q for q in json.loads(current[0])['questions']} if current else {}
            added=[q['id'] for q in payload['questions'] if q['id'] not in old]
            changed=[q['id'] for q in payload['questions'] if q['id'] in old and q!=old[q['id']]]
            removed=[key for key in old if key not in {q['id'] for q in payload['questions']}]
            c.execute('INSERT OR IGNORE INTO question_bank_versions VALUES(?,?,?,?,?,?,?)',
                (vid,name[:120],fingerprint,hashlib.sha256(content).hexdigest(),raw,int(time.time()*1000),author))
        return {**self.info(vid),'added':added,'changed':changed,'removed':removed,'warnings':payload['warnings']}

    def info(self,vid):
        with self.db.connect() as c:
            row=c.execute('SELECT * FROM question_bank_versions WHERE id=?',(vid,)).fetchone()
            active_table='relaxation_bank_active' if vid.startswith('relax-bank-') else 'question_bank_active'
            active=c.execute(f'SELECT version_id FROM {active_table} WHERE singleton=1').fetchone()
        if not row:raise ApiError(404,'BANK_NOT_FOUND','题库版本不存在')
        payload=json.loads(row['payload_json']);qs=payload['questions']
        if payload.get('kind')=='relaxation':
            return {'id':vid,'kind':'relaxation','name':row['name'],'count':len(qs),
                'activeCount':sum(q['status']=='active' for q in qs),'heldCount':sum(q['status']!='active' for q in qs),
                'createdAt':row['created_ms'],'active':bool(active and active[0]==vid),'sourceHash':row['source_hash'],
                'categories':{category:sum(q['category']==category for q in qs) for category in sorted({q['category'] for q in qs})},
                'declaredCount':payload['declaredCount'],'missingIds':payload['missingIds'],'warnings':payload['warnings']}
        return {'id':vid,'name':row['name'],'count':len(qs),'createdAt':row['created_ms'],'active':bool(active and active[0]==vid),
            'kind':'assessment',
            'sourceHash':row['source_hash'],'pools':{p:sum(q['pool']==p for q in qs) for p in WEIGHTS},
            'dimensions':{d:sum(q['dimension']==d for q in qs) for d in DIMENSIONS}}

    def versions(self):
        with self.db.connect() as c:ids=[r[0] for r in c.execute('SELECT id FROM question_bank_versions ORDER BY created_ms DESC LIMIT 100')]
        return [self.info(vid) for vid in ids]

    def activate(self,vid,author):
        info=self.info(vid)
        active_table='relaxation_bank_active' if info['kind']=='relaxation' else 'question_bank_active'
        with self.db.connect() as c:
            c.execute(f'INSERT INTO {active_table} VALUES(1,?) ON CONFLICT(singleton) DO UPDATE SET version_id=excluded.version_id',(vid,))
            c.execute('INSERT INTO admin_audit(username,action,target,created) VALUES(?,?,?,?)',(author,'activate_question_bank',vid,int(time.time())))
        return self.info(vid)

    def get_payload(self,vid):
        with self.db.connect() as c:row=c.execute('SELECT payload_json FROM question_bank_versions WHERE id=?',(vid,)).fetchone()
        if not row:raise ApiError(404,'BANK_NOT_FOUND','题库版本不存在')
        return json.loads(row[0])

    def _assignment(self,c,user_id,mode):
        assigned=c.execute('SELECT * FROM question_bank_assignments WHERE user_id=? AND mode=?',(user_id,mode)).fetchone()
        if assigned:return dict(assigned)
        # Adopt the latest pre-policy form without changing its frozen questions.
        # Existing accounts keep their last set and receive the new six-refresh
        # allowance. Only explicit refreshes under this policy consume it.
        row=c.execute('SELECT id FROM question_bank_forms WHERE user_id=? AND mode=? ORDER BY created_ms DESC,rowid DESC LIMIT 1',(user_id,mode)).fetchone()
        if not row:return None
        return {'user_id':user_id,'mode':mode,'source_form_id':row['id'],'refresh_used':0}

    def _quota(self,c,user_id):
        assignments={mode:self._assignment(c,user_id,mode) for mode in ('standard','challenge')}
        used=sum(item['refresh_used'] for item in assignments.values() if item)
        return {'limit':QUESTION_REFRESH_LIMIT,'used':used,'remaining':max(0,QUESTION_REFRESH_LIMIT-used),
            'scope':'account','period':'lifetime',
            'modes':{mode:{'hasOriginal':item is not None,'questionSetId':item['source_form_id'] if item else None} for mode,item in assignments.items()}}

    def quota(self,user_id):
        with self.db.connect() as c:
            c.execute('BEGIN')
            return self._quota(c,user_id)

    def issue_form(self,user_id,mode,request_id,include_relaxation=False,shuffle_relaxation=False,refresh=False):
        if mode not in ('standard','challenge') or not ID_RE.fullmatch(request_id):raise ApiError(400,'BANK_FORM_INVALID','出题请求格式不正确')
        with self.db.connect() as c:
            c.execute('BEGIN IMMEDIATE')
            old=c.execute('SELECT * FROM question_bank_forms WHERE user_id=? AND request_id=?',(user_id,request_id)).fetchone()
            if old:
                if old['mode']!=mode:raise ApiError(409,'BANK_FORM_MODE_MISMATCH','此出题请求已绑定另一个模式')
                if bool(old['refresh_requested'])!=refresh:raise ApiError(409,'BANK_FORM_INTENT_MISMATCH','此请求已绑定原题或换题操作，请继续原操作')
                form_id=old['id']
            else:
                for item_mode in ('standard','challenge'):
                    existing=self._assignment(c,user_id,item_mode)
                    if existing:c.execute('INSERT OR IGNORE INTO question_bank_assignments VALUES(?,?,?,?)',
                        (user_id,item_mode,existing['source_form_id'],existing['refresh_used']))
                assigned=self._assignment(c,user_id,mode)
                if refresh and not assigned:raise ApiError(409,'BANK_ORIGINAL_REQUIRED','请先领取该模式的首次题组，首次领取不消耗换题次数')
                if refresh and self._quota(c,user_id)['remaining']==0:
                    raise ApiError(409,'BANK_REFRESH_LIMIT_REACHED','6次换题机会已用完，你仍可使用原题测评')
                form_id=uuid.uuid4().hex
                if assigned and not refresh:
                    source=c.execute('SELECT * FROM question_bank_forms WHERE id=? AND user_id=?',(assigned['source_form_id'],user_id)).fetchone()
                    version_id=source['version_id'];chosen=json.loads(source['question_ids_json'])
                    include_relaxation=bool(source['include_relaxation']);shuffle_relaxation=bool(source['shuffle_relaxation'])
                    payload=c.execute('SELECT payload_json FROM question_bank_versions WHERE id=?',(version_id,)).fetchone()[0]
                    lookup={q['id']:q for q in json.loads(payload)['questions']}
                    composition=self._composition(c,user_id,source['id'],[lookup[qid] for qid in chosen],include_relaxation,shuffle_relaxation)
                    source_id=source['id']
                else:
                    version=c.execute('SELECT v.id,v.payload_json FROM question_bank_active a JOIN question_bank_versions v ON a.version_id=v.id').fetchone()
                    if not version:raise ApiError(503,'BANK_NOT_READY','题库正在准备，请稍后重试',retryable=True)
                    version_id=version['id'];questions=json.loads(version['payload_json'])['questions'];chosen=[]
                    rng=random.SystemRandom()
                    previous=set();previous_version=None
                    if assigned:
                        old_set=c.execute('SELECT question_ids_json,version_id FROM question_bank_forms WHERE id=?',(assigned['source_form_id'],)).fetchone()
                        previous=set(json.loads(old_set['question_ids_json']));previous_version=old_set['version_id']
                    for d in DIMENSIONS:
                        for pool,n in [('anchor',2),(mode,1)]:
                            group=[q['id'] for q in questions if q['dimension']==d and q['pool']==pool and q['status']!='retired']
                            fresh=[qid for qid in group if qid not in previous]
                            selected=rng.sample(fresh,min(n,len(fresh)))
                            selected+=rng.sample([qid for qid in group if qid not in selected],n-len(selected))
                            chosen.extend(selected)
                    if refresh and version_id==previous_version and set(chosen)==previous:
                        raise ApiError(409,'BANK_NO_NEW_QUESTIONS','当前题库暂无可更换的新题，本次不消耗次数')
                    rng.shuffle(chosen);composition=None;source_id=form_id
                c.execute('INSERT INTO question_bank_forms(id,user_id,request_id,version_id,mode,question_ids_json,created_ms,include_relaxation,shuffle_relaxation,source_form_id,refresh_requested) VALUES(?,?,?,?,?,?,?,?,?,?,?)',
                    (form_id,user_id,request_id,version_id,mode,encode(chosen),int(time.time()*1000),int(include_relaxation),int(include_relaxation and shuffle_relaxation),source_id,int(refresh)))
                if composition is None:
                    lookup={q['id']:q for q in questions}
                    self._composition(c,user_id,form_id,[lookup[qid] for qid in chosen],include_relaxation,shuffle_relaxation)
                else:c.execute('INSERT INTO question_bank_compositions VALUES(?,?)',(form_id,encode(composition)))
                if assigned and refresh:
                    c.execute('UPDATE question_bank_assignments SET source_form_id=?,refresh_used=refresh_used+1 WHERE user_id=? AND mode=?',(form_id,user_id,mode))
                elif not assigned:c.execute('INSERT INTO question_bank_assignments VALUES(?,?,?,0)',(user_id,mode,form_id))
        return self.form(user_id,form_id)

    def form(self,user_id,form_id):
        with self.db.connect() as c:row=c.execute('SELECT * FROM question_bank_forms WHERE id=? AND user_id=?',(form_id,user_id)).fetchone()
        if not row:raise ApiError(404,'BANK_FORM_NOT_FOUND','题表不存在或不属于当前账号')
        lookup={q['id']:q for q in self.get_payload(row['version_id'])['questions']}
        ids=json.loads(row['question_ids_json']); qs=[lookup[qid] for qid in ids]
        return {'id':form_id,'version':row['version_id'],'mode':row['mode'],'questions':qs,
            'questionSetId':row['source_form_id'] or row['id'],
            'composition':self.composition(user_id,form_id,qs,bool(row['include_relaxation']),bool(row['shuffle_relaxation']))}

    def composition(self,user_id,form_id,questions,include_relaxation=False,shuffle_relaxation=False):
        with self.db.connect() as c:
            c.execute('BEGIN IMMEDIATE')
            return self._composition(c,user_id,form_id,questions,include_relaxation,shuffle_relaxation)

    def _composition(self,c,user_id,form_id,questions,include_relaxation=False,shuffle_relaxation=False):
        from .question_composer import compose_round
        old=c.execute('SELECT payload_json FROM question_bank_compositions WHERE form_id=?',(form_id,)).fetchone()
        if old:return json.loads(old[0])
        previous=c.execute('SELECT p.payload_json FROM question_bank_compositions p JOIN question_bank_forms f ON f.id=p.form_id WHERE f.user_id=? ORDER BY f.created_ms DESC LIMIT 5',(user_id,)).fetchall()
        recent=[q['id'] for r in previous for q in json.loads(r[0])['icebreakers']]
        relaxation=c.execute('SELECT v.id,v.payload_json FROM relaxation_bank_active a JOIN question_bank_versions v ON v.id=a.version_id').fetchone() if include_relaxation else None
        result=compose_round(questions,recent,json.loads(relaxation['payload_json'])['questions'] if relaxation else None,shuffle_relaxation=shuffle_relaxation)
        if relaxation:result['relaxationBankVersion']=relaxation['id']
        c.execute('INSERT INTO question_bank_compositions VALUES(?,?)',(form_id,encode(result)))
        return result

    def assessment_snapshot(self,user_id,raw):
        from .paid_contract import assessment_snapshot
        if not isinstance(raw,dict) or not isinstance(raw.get('formId'),str):
            raise ApiError(400,'ASSESSMENT_INVALID','测评快照格式不完整')
        issued=self.form(user_id,raw.get('formId',''))
        if raw.get('formVersion')!=issued['version'] or raw.get('mode')!=issued['mode']:
            raise ApiError(400,'BANK_FORM_MISMATCH','答案与服务器题表不一致')
        qs=issued['questions']
        form={'id':issued['id'],'anchors':{d:[q['id'] for q in qs if q['dimension']==d and q['pool']=='anchor'] for d in DIMENSIONS},
            issued['mode']:{d:next(q['id'] for q in qs if q['dimension']==d and q['pool']==issued['mode']) for d in DIMENSIONS}}
        return assessment_snapshot(raw,catalog={'forms':[form],'items':qs,'modes':[issued['mode']]})

    def save_result(self,user_id,raw):
        snapshot=self.assessment_snapshot(user_id,raw)
        canonical=encode(snapshot)
        with self.db.connect() as c:
            c.execute('BEGIN IMMEDIATE')
            existing=c.execute('SELECT snapshot_json FROM question_bank_results WHERE form_id=?',(snapshot['formId'],)).fetchone()
            if existing and existing[0]!=canonical:
                raise ApiError(409,'ASSESSMENT_ALREADY_SUBMITTED','本轮结果已保存，重新测评需领取新题表')
            c.execute('INSERT OR IGNORE INTO question_bank_results VALUES(?,?,?)',(snapshot['formId'],canonical,int(time.time()*1000)))
        return {'formId':snapshot['formId'],'saved':True,'dimensions':snapshot['dimensions']}
